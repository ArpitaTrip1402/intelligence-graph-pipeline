import json
from openpyxl import Workbook

from src.pipeline.entity_resolution import create_mapping_record


def load_research_papers():

    with open(
        "data/research_papers.json",
        "r",
        encoding="utf-8"
    ) as file:

        return json.load(file)


def create_dataset():

    wb = Workbook()

    default_sheet = wb.active
    wb.remove(default_sheet)

    startups = wb.create_sheet("Startups")
    products = wb.create_sheet("Products")
    papers = wb.create_sheet("Research Papers")
    jobs = wb.create_sheet("Jobs")
    news = wb.create_sheet("News")
    entity_log = wb.create_sheet("Entity Mapping Log")

    # Startups
    startups.append([
        "name",
        "website",
        "source_url"
    ])

    # Products
    products.append([
        "name",
        "company",
        "website",
        "source_url"
    ])

    # Research Papers
    papers_data = load_research_papers()

    if papers_data:

        papers.append(list(papers_data[0].keys()))

        for paper in papers_data:
            papers.append(list(paper.values()))

    # Jobs
    jobs.append([
        "title",
        "company",
        "location",
        "url",
        "published_date"
    ])

    # News
    news.append([
        "title",
        "source",
        "url",
        "published_date"
    ])

    # Entity Mapping Log
    entity_log.append([
        "raw_entity",
        "canonical_entity",
        "entity_type",
        "reason"
    ])

    entities = [
        ("Open AI", "COMPANY"),
        ("OpenAI Inc.", "COMPANY"),
        ("DeepMind", "COMPANY"),
        ("Google DeepMind", "COMPANY"),
        ("Microsoft Corp", "COMPANY"),
        ("Meta Platforms", "COMPANY")
    ]

    for raw_entity, entity_type in entities:

        record = create_mapping_record(
            raw_entity,
            entity_type
        )

        entity_log.append([
            record["raw_entity"],
            record["canonical_entity"],
            record["entity_type"],
            record["reason"]
        ])

    wb.save(
        "data/intelligence_graph_dataset.xlsx"
    )

    print("Dataset created successfully.")
    print(
        "File: data/intelligence_graph_dataset.xlsx"
    )


if __name__ == "__main__":
    create_dataset()